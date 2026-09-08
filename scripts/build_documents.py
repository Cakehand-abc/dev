"""Build ten project documents from the selected retained System Design template.
Run with the Codex bundled Python. Render all outputs with render_docx.py afterwards.
"""
from pathlib import Path
from copy import deepcopy
from datetime import datetime
import json,re,hashlib
from xml.etree import ElementTree as ET
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.table import Table

ROOT=Path(__file__).resolve().parents[1]
REF=Path('C:/Users/29569/.codex/plugins/cache/openai-curated-remote/openai-templates/0.1.1/skills/artifact-template-system-design/assets/reference.docx')
OUT=ROOT/'docs/deliverables';OUT.mkdir(parents=True,exist_ok=True)
SOURCE=ROOT/'docs/source';SOURCE.mkdir(exist_ok=True)
DATE='2026年9月8日';PROJECT='智慧医养大数据决策分析系统'
ref=Document(REF);table_patterns={2:deepcopy(ref.tables[2]._tbl),3:deepcopy(ref.tables[5]._tbl),4:deepcopy(ref.tables[4]._tbl)}
perf=json.loads((ROOT/'docs/evidence/performance.json').read_text(encoding='utf-8-sig'))
tests=[]
for f in (ROOT/'target/surefire-reports').glob('TEST-*.xml'):
 t=ET.parse(f).getroot();tests += [{'name':x.attrib['name'],'class':x.attrib['classname'],'seconds':float(x.attrib['time']),'passed':x.find('failure') is None and x.find('error') is None} for x in t.findall('testcase')]
(ROOT/'docs/evidence/maven-tests.json').write_text(json.dumps({'runDate':DATE,'java':'25.0.1','compilerRelease':17,'springBoot':'3.5.16','tests':tests},ensure_ascii=False,indent=2),encoding='utf-8')

def P(s):return ('p',s)
def H(s):return ('h',s)
def T(headers,rows):return ('t',headers,rows)
def IMG(file,caption):return ('img',file,caption)

ROLES=[['管理员 ADMIN','全部查询、业务维护、账号管理、demo 数据入口'],['业务人员 OPERATOR','全部业务查询、档案录入、设备绑定、围栏配置及事件处理'],['分析员 ANALYST','业务查询、统计与轨迹查看；无写入和账号管理权限']]
STACK=[['Java 与构建','Java 17 编译目标；Maven Wrapper 3.9.16'],['后端框架','Spring Boot 3.5.16；Spring MVC；Spring Security；Spring 事务'],['数据访问','MyBatis-Plus 3.5.17 BaseMapper；MyBatis XML 聚合统计'],['数据库','MySQL 8.0.16 及以上 8.0 系列；当前实测 8.0.46；H2 仅测试'],['前端','Vue 3、Vite 7、ECharts 6；npm 锁定依赖'],['测试','Boot BOM 管理 JUnit Jupiter 5.12.2 与 Platform 1.12.2']]
FR=[
 ['FR01','登录与退出','凭据正确进入首页；未认证 API 返回 401；退出后会话失效'],
 ['FR02','首页总览','地区和日期筛选；人口、健康、随访、设备与围栏概览；入口跳转'],
 ['FR03','人口统计','新增、编辑、归档和分页查询；年龄、性别、地区分组；未知值单列'],
 ['FR04','健康检测','至少一项指标；时间校验；规则快照；检测次数与人数分开；个体趋势'],
 ['FR05','医生随访','创建计划、完成或取消；计划最多一次完成；零分母显示无计划'],
 ['FR06','腕表分布','设备维护、绑定和解绑；心跳状态、最新有效定位与无定位提示'],
 ['FR07','电子围栏','圆形围栏及成员管理；越界一次告警；返回关闭；人工处理留痕'],
 ['FR08','轨迹回放','按老人查询 24 小时内轨迹；最多 5000 点；播放暂停倍速；缺口标记']]
RULES=[
 ['时间','API 带时区偏移，入库统一 Asia/Shanghai，精确到微秒。筛选采用 [start,end)，结束日期在前端换成次日零点。'],
 ['人口','仅统计 ACTIVE 档案，按查询当天周岁与当前地区分组；不提供历史人口快照。年龄未知、性别未知分别计入。'],
 ['健康','次数按记录数，人数按 elder_id 去重。至少一项指标超演示阈值即为异常；缺测保留 null；保存规则版本和快照。'],
 ['随访','分母为期间到期且截至截止时未取消的计划；分子为截止时已完成计划。截止为 end 与当前时间较早者。零计划返回 null 完成率。'],
 ['设备','最近心跳不超过 5 分钟为在线，无心跳为未知；停用单列。静态演示数据随时间转离线属于正常现象。'],
 ['围栏','WGS84 球面距离大于半径判越界；首点在外也告警；连续在外不重复；返回后再离开产生新事件。'],
 ['乱序与绑定','以设备和 eventId 去重，相同事件不同内容返回 409。历史定位按发生时间查找绑定，不归到当前新绑定人。'],
 ['轨迹','按 recorded_at、id 稳定升序。相邻点间隔超过 10 分钟显示缺口；不补点、不推测道路。']]
LIMITS='本版用于单单位教学演示，使用虚构老人和设备数据。坐标画布没有道路底图；未接真实腕表、短信、第三方地图、Agent、分布式计算或多租户。健康规则为教学参数，不承担临床诊断用途。'

docs=[]
def add(number,title,owner,scope,pages):docs.append({'number':number,'title':title,'owner':owner,'scope':scope,'pages':pages})
add('01','软件需求规范','A 项目经理，B C D 参与','定义八项功能、权限、统计口径及验收边界',[
 [H('1 项目目的与范围'),P('本系统面向医养服务管理与分析人员，将老人档案、健康检测、医生随访及设备定位整合到同一浏览器工作台，支持服务情况查询和统计分析。项目背景采用用户提供的医养结合示范项目建设要求，不宣称已经获得示范项目认定。'),P(LIMITS),H('2 用户与权限'),T(['系统角色','允许操作'],ROLES),H('3 工程约束'),T(['项目','实现基线'],STACK),P('按团队要求使用 Spring Boot 3，避免引入 Boot 4 技术线。IDEA 2024 及更早版本的完整版本号仍待团队补充，兼容性以实际运行记录为准。')],
 [H('4 功能需求'),T(['编号','功能','验收条件'],FR),H('5 配套管理'),P('管理员维护账号并重置密码；业务人员维护医生。老人归档会关闭绑定与监测、取消未完成随访计划，保留历史记录。修改采用版本校验，过期版本返回冲突并要求刷新。'),P('页面提供加载、空数据和失败状态，列表支持分页。电话在响应中脱敏；系统不采集真实身份证号码。')],
 [H('6 数据与统计口径'),T(['主题','规则'],RULES[:4]),H('7 设备与空间规则'),T(['主题','规则'],RULES[4:]),P('圆形围栏半径为 50 至 5000 米。定位时间允许最多 60 秒时钟偏差，超范围坐标和未来异常时间拒收；旧定位允许存档，但不回退监测状态。')],
 [H('8 非功能需求'),T(['编号','要求','验证方法'],[['NFR01','未登录不可读取业务数据，写接口必须通过 CSRF 与角色检查','MockMvc 权限测试及实际会话冒烟'],['NFR02','10 万健康与 10 万定位数据，20 并发 60 秒，常规查询 P95 ≤ 2 秒、错误率 < 1%','隔离 MySQL 压力测试，记录端点和样本'],['NFR03','事务保证完成随访、绑定与告警状态一致','业务断言、唯一约束和外键'],['NFR04','同源打包部署，后端独立 Maven 测试','构建脚本、JAR 启动、前端构建'],['NFR05','1440 桌面与 390 像素窄屏可访问；表格局部滚动','浏览器检查及逐页截图']]),H('9 验收与变更'),P('以八项功能演示、自动测试、接口检查和交付文件核对作为课程版验收依据。当前验证结果见测试日志；实际 IDEA 直接运行、团队 JDK 17、真实设备和公开生产部署不记为已通过。'),P('新增真实道路底图、外部 Agent、付费接口、真实个人数据、清空数据库或公开部署属于范围或环境变化，先评估影响并由用户确认。成员姓名、学校格式与截止日期由团队补充。'),H('10 追踪关系'),T(['需求','实现位置','验证'],[['FR01','auth 与 SecurityConfig','TC01 至 TC06'],['FR02 至 FR05','CareService、StatisticsService、App.vue','TC07 至 TC13、TC21、HTTP 与浏览器'],['FR06 至 FR08','GeoService、GeoCanvas.vue','TC14 至 TC20、TC22、浏览器']])]
])

schema=(ROOT/'database/01-schema.sql').read_text(encoding='utf-8')
def split_sql(s):
 parts=[];depth=0;start=0;quote=False
 for i,ch in enumerate(s):
  if ch=="'":quote=not quote
  if not quote:
   if ch=='(':depth+=1
   elif ch==')':depth-=1
   elif ch==',' and depth==0:parts.append(s[start:i].strip());start=i+1
 parts.append(s[start:].strip());return parts
tables=[]
for m in re.finditer(r'CREATE TABLE (\w+) \((.*?)\);',schema,re.S):
 cols=[];cons=[]
 for s in split_sql(m[2]):
  c=re.match(r'(\w+)\s+(BIGINT|INT|VARCHAR\(\d+\)|BOOLEAN|DATE\b|DATETIME\(6\)|DOUBLE|DECIMAL\(\d+,\d+\)|TEXT)(.*)',s,re.S)
  if c:cols.append((c[1],c[2],c[3].strip()))
  else:cons.append(' '.join(s.split()))
 tables.append((m[1],cols,cons))
names={'sys_user':'系统账号','region':'地区','elder':'老人档案','doctor':'医生','health_rule':'教学阈值规则','health_record':'健康检测记录','followup_plan':'随访计划','followup_record':'随访结果','watch_device':'腕表设备','device_binding':'历史绑定','location_point':'定位点','geofence':'圆形围栏','geofence_member':'围栏成员','geofence_state':'围栏监测状态','geofence_alert':'越界事件','audit_log':'操作审计'}
labels=dict(x.split('=',1) for x in 'id=自增主键|username=登录名|password_hash=BCrypt 密文|display_name=显示名|role=ADMIN OPERATOR ANALYST|enabled=是否启用|auth_version=认证版本|code=业务编号|name=名称|parent_id=父地区主键|gender=MALE FEMALE UNKNOWN|birth_date=出生日期 可未知|region_id=地区主键|phone=电话 响应脱敏|status=业务状态|version=并发版本|department=科室|metric=检测指标名|lower_bound=下界 包含|upper_bound=上界 包含|unit=单位|elder_id=老人主键|measured_at=检测时间|systolic=收缩压 mmHg|diastolic=舒张压 mmHg|heart_rate=心率 次每分|oxygen=血氧百分比|temperature=体温 摄氏度|abnormal=是否超教学阈值|rule_snapshot=判定规则 JSON 快照|source=数据来源|event_id=外部事件编号|doctor_id=医生主键|due_at=计划到期时间|canceled_at=取消时间|cancel_reason=取消原因|plan_id=随访计划主键 唯一|completed_at=完成时间|content=随访内容|result=结果摘要|serial_no=设备序列号|model=设备型号|last_seen_at=最后心跳|device_id=设备主键|bound_at=绑定开始时间|unbound_at=解绑时间|active_device_id=有效设备唯一占位|active_elder_id=有效老人唯一占位|binding_id=发生时绑定主键|longitude=WGS84 经度|latitude=WGS84 纬度|recorded_at=设备记录时间|received_at=服务端接收时间|center_lon=围栏中心经度|center_lat=围栏中心纬度|radius_m=半径 米|fence_id=围栏主键|member_id=围栏成员主键|last_point_id=最后参与监测的定位|last_recorded_at=监测事件时间|inside=是否在围栏内 可未知|active_alert_id=当前未结束越界事件|exit_point_id=触发定位点|triggered_at=越界触发时间|returned_at=实际返回时间|closed_at=监测关闭时间|close_reason=关闭原因|handled_by=处理账号|handled_at=处理时间|handling_note=处理说明|geometry_snapshot=围栏几何 JSON 快照|user_id=操作账号 可系统执行|action=动作编码|target_type=对象类型|target_id=对象主键|occurred_at=操作时间'.split('|'))
dbpages=[[H('1 数据库概览'),P('数据库 smart_care 使用 MySQL 8.0，包含 16 张表。database/01-schema.sql 为结构权威来源；首次在专用空库执行，重复执行会失败，不包含删除或覆盖已有数据库的逻辑。'),T(['分组','表与职责'],[['身份和档案','sys_user、region、elder、doctor'],['健康与随访','health_rule、health_record、followup_plan、followup_record'],['设备与定位','watch_device、device_binding、location_point'],['围栏与审计','geofence、geofence_member、geofence_state、geofence_alert、audit_log']]),H('2 关系与生命周期'),P('地区一对多老人；老人一对多健康记录和随访计划；医生一对多计划；计划与完成记录为一对零或一。设备和老人通过历史绑定建立关系，当前绑定分别唯一。定位点同时保留设备、绑定及老人外键。'),P('围栏通过成员表关联老人。每名成员最多一条监测状态；状态指向当前活动事件，事件指向越界定位与处理账号。先建立状态和事件表，最后追加 active_alert_id 外键以处理循环引用。'),P('业务主键为 BIGINT，JSON 响应中的 Long 以字符串输出，避免 JavaScript 精度损失。业务 DATETIME(6) 保存上海本地时间；地区、规则与演示数据初始化由应用控制。')]]
for i,(name,cols,cons) in enumerate(tables):
 rows=[]
 for field,typ,extra in cols:
  null='否' if 'NOT NULL' in extra or 'PRIMARY KEY' in extra else '是'
  flags='主键 自增' if 'PRIMARY KEY' in extra else ('唯一 ' if 'UNIQUE' in extra else '')
  default=re.search(r'DEFAULT\s+(\S+)',extra)
  if default:flags+='默认 '+default[1]
  rows.append([field,typ,('可空 ' if null=='是' else '必填 ')+flags,({('health_rule','version'):'规则版本',('audit_log','result'):'执行结果码'}.get((name,field),labels.get(field,field)))])
 page=[H(f'{i+3} {names[name]} {name}'),T(['字段','类型','约束与默认','含义'],rows)]
 if cons:page.append(P('表级约束：'+'；'.join(cons)))
 indexes=re.findall(r'CREATE INDEX ([^;]+) ON '+name+r'\(([^;]+)\);',schema)
 if indexes:page.append(P('索引：'+'；'.join(a+' ('+b+')' for a,b in indexes)))
 dbpages.append(page)
dbpages.append([H('19 索引与完整性设计'),T(['约束','作用'],[['source 与 event_id 唯一','健康重复录入产生冲突，客户端可识别重试'],['device_id 与 event_id 唯一','定位重传不重复入库'],['active_device_id 和 active_elder_id 分别唯一','一个设备和老人同时最多一条有效绑定，历史解绑后置空'],['followup_record.plan_id 唯一','防止同一随访计划重复完成'],['围栏成员及状态唯一','防止同一关联存在多条监测状态']]),P('归档以状态变更保留历史，不级联删除。定位、告警和绑定属于历史事实；修改围栏保存新的几何状态并关闭原监测事件，旧事件保留几何快照。'),H('20 初始化 备份与迁移'),P('先创建专用数据库，再执行表结构，再由 Bootstrap 建管理员、地区与教学规则。demo 模式且 APP_SEED_DEMO=true 时，空老人表触发演示种子；非空库不会重复生成。独立 SQL 种子见 database/02-demo-data.sql，与 Java 种子二选一。'),P('结构调整应新增迁移脚本并先备份验证。当前项目没有自动 Flyway 迁移，也没有自动备份任务。导出与恢复命令见部署说明；生产恢复必须先在独立库验证，不覆盖现有业务库。')])
add('02','数据库设计说明书','C 后端兼数据工程师','说明实际 16 张表 全字段 索引及数据生命周期',dbpages)

perfrows=[[e['path'],str(e['samples']),str(e['p95Milliseconds'])+' ms',str(e['errors'])] for e in perf['endpoints']]
add('03','测试日志','A 测试负责人汇总','记录真实构建 联调 缺陷回归及未执行项',[
 [H('1 测试结论'),P('截至 2026年9月8日，本地课程演示版的 Maven 测试、MySQL 接口冒烟和 Chrome 页面检查通过。性能验证满足本次约定的四类查询目标。此结论不等同真实设备或公开生产环境验收。'),T(['项目','结果','证据'],[['Maven JUnit','23 项通过 0 失败 0 错误','docs/evidence/maven-tests.json'],['MySQL 实际 HTTP','16 项通过','docs/evidence/http-smoke.json'],['Chrome 页面检查','12 项通过 无页面脚本异常','docs/evidence/browser-qa.json'],['负载测试','2664 请求 0 错误 四端点 P95 < 2 秒','docs/evidence/performance.json']]),H('2 执行环境'),T(['项目','实测'],[['操作系统','Windows 11 10.0.26100'],['CPU','Intel Core i7-13650HX 14 核 20 逻辑处理器'],['Java','Zulu 25.0.1 编译 release 17；未冒充 JDK 17 实测'],['数据库','隔离 MySQL 8.0.46 端口 13316，16 张表'],['应用和浏览器','JAR 端口 18080；Chrome 152.0.7977.83'],['构建','Maven Wrapper 3.9.16；Vite 7.3.6']])],
 [H('3 实际测试过程与修复'),T(['记录','问题与原因','修复及回归'],[['D01','统计 value 别名与 H2 保留字冲突','为别名加反引号；统计测试与 MySQL 冒烟通过'],['D02','先添加 ORDER BY 再调用 selectCount 导致计数 SQL 不兼容','计数后再添加排序；分页与轨迹测试通过'],['D03','Java 纳秒与数据库微秒精度不同导致同事件比较不一致','输入与当前时间统一截断微秒；重传与乱序测试通过'],['D04','窄屏饼图标签越出图表边界','取消外侧标签，图例显示名称和值；重新截图检查'],['D05','Windows 运行中的演示 JAR 阻止重新打包改名','停止本任务进程后打包成功；部署说明写明先停止'],['D06','窄屏设备详情被样式隐藏','详情改为在地图下方显示；前端构建与窄屏复核']]),P('以上为开发过程中的实际缺陷和环境问题，不补造发现人的姓名、工时或线下会议。测试执行与页面检查由开发自动化完成，A 负责后续人工复核与签收。'),H('4 本次构建记录'),P('2026年9月8日 00时26分，22 项 BusinessTests 与 1 项应用上下文测试通过。随后打包阶段遇到运行中 JAR 占用；停止演示进程后，使用已通过测试的编译结果再次 package 成功。最终交付按部署说明可重新构建。')],
 [H('5 性能实测'),P('2026年9月8日 00时28分开始。原库 1080 条健康、648 条定位；各加入 100000 条带本次唯一前缀的模拟记录，测试时分别为 101080 和 100648 条。20 个工作线程持续约 60.371 秒，默认近 30 天范围、首页 20 条、无地区过滤。'),T(['端点','样本数','P95','错误数'],perfrows),P('端点均满足 P95 ≤ 2000 ms、错误率 < 1%。结果仅代表本机、本数据分布与这四类只读端点，不代表并发写入、全量轨迹导出或分布式规模。测试脚本在 finally 中删除本次专属记录，已核对恢复为 1080 和 648 条。'),H('6 未执行和待签收'),T(['检查','当前状态','后续负责人'],[['JDK 17 真正运行','未执行；当前机器测试 JVM 为 25','B 配置团队 JDK 17 后运行 mvnw test'],['IDEA 内置 JUnit 运行','未执行；完整 IDEA 版本待补充','B 记录版本和右键运行结果'],['真实腕表及道路地图','不在本版范围','C 在新增范围获批后对接'],['正式部署 恢复演练','未执行公开部署及备份恢复','A C 确认目标环境后执行']]),P('人工复核人：待团队填写。验收日期与签字：待团队实际验收后填写。')]
])

add('04','系统架构设计说明书','B C 后端工程师','解释系统分层 关键边界 部署和技术取舍',[
 [H('1 架构结论'),P('系统采用 Vue 单页前端与 Spring Boot 模块化单体。构建后将 Vue 静态资源纳入可执行 JAR，由同一服务提供页面与 /api/v1 接口。MySQL 保存全部业务事实，服务端 Session 保存认证状态。该方案减少四人小组的部署与联调复杂度。'),T(['层级','组件','责任'],[['展示层','App.vue Chart.vue GeoCanvas.vue','页面筛选、交互表单、统计图与坐标回放'],['访问边界','Spring Security 与 AuthController','会话、CSRF、角色与账号版本校验'],['接口层','CareController DemoController Api','接收参数、统一响应、演示环境边界'],['业务层','CareService GeoService StatisticsService','事务、不变量、统计口径与空间规则'],['持久层','BaseMapper 与 StatisticsMapper XML','CRUD、行锁、聚合查询、索引访问'],['存储层','MySQL 16 张表','外键、唯一约束和历史事实']]),H('2 调用链'),P('浏览器先获取 CSRF，再提交登录。服务端认证后轮换 Session 并返回角色，前端重新获取 CSRF。业务请求经过账号有效性检查、权限检查、输入白名单与业务事务后访问 Mapper；返回统一 JSON，再由 Vue 更新页面。'),P('只读统计同时使用 MyBatis XML；普通增删改查使用 MyBatis-Plus。Spring MVC、Spring DI、Spring 事务均在真实业务调用链中使用。')],
 [H('3 模块与数据边界'),T(['模块','核心对象','不变量'],[['人口与健康','老人、地区、检测、规则','归档老人不进人口总数；检测保存规则快照'],['随访','医生、计划、完成记录','计划最多完成一次；取消和完成终态互斥'],['设备','设备、绑定、定位','有效绑定双侧唯一；历史定位归属不随换绑改变'],['围栏','成员、状态、事件','持续越界一条事件；人工处理与实际返回分离'],['账号与审计','账号、认证版本、审计','保留最后管理员；角色变更使旧会话失效']]),H('4 一致性与并发'),P('关键业务方法使用 @Transactional，异常触发回滚。绑定依赖行锁和数据库唯一约束；随访完成锁定计划，再写结果与状态。编辑通过 version 比较拒绝过期提交。定位按事件时间和主键推进状态，旧点只保留历史。'),P('围栏、绑定和定位采用统一锁顺序。当前 GeoService 使用单个账号行作为监测串行锁，再按围栏编号锁定，优先保证教学版一致性；这会限制并发写入吞吐，未来应改为设备分区锁或消息分区消费，并补充并发测试。')],
 [H('5 安全和故障处理'),T(['机制','已实现行为'],[['认证','BCrypt 密码；30 分钟 Session；注销使会话失效'],['权限','后端约束 ADMIN OPERATOR ANALYST；只读角色不可写'],['跨站请求','写接口验证 CSRF；同源访问；Cookie HttpOnly 与 SameSite=Lax'],['输入','字段白名单、数值时间范围、分页上限和未知字段拒收'],['数据','电话脱敏；不收集身份证；演示凭据放本机私有文件'],['错误','400 参数；401 未登录；403 禁止；404 不存在；409 冲突；500 未预期异常']]),P('数据库不可用时请求失败，不展示伪造数据；前端显示失败状态。系统没有自动离线缓存、熔断集群或故障转移。审计表记录成功的关键业务操作，不是全量访问日志或独立安全监控产品。'),H('6 部署结构'),P('开发态：Vite 开发服务器代理 /api 到后端 8080。交付态：浏览器同源访问 JAR，JAR 连接 MySQL。当前演示因本机 8080 已占用，使用 18080；隔离数据库使用 13316。prod 配置启用 Secure Cookie，需要 HTTPS 入口。'),P('本版默认仅监听 127.0.0.1，不自动配置公网端口、证书或防火墙。数据库需使用上海时区，避免 CURRENT_DATE 与应用业务日期不一致。')],
 [H('7 技术决策'),T(['决策','原因','影响'],[['Spring Boot 3.5.16','团队明确要求 Boot 3 与 JUnit 5 技术线','BOM 管理版本；IDEA 环境还需实际确认'],['单体加 MySQL','四人可共同构建演示、定位事务问题','没有独立微服务部署和跨节点 Session'],['离线坐标画布','不依赖密钥，围栏和轨迹可完整演示','无道路、卫星底图与导航能力'],['不设 Agent 工程师','当前功能均可由确定性查询与规则实现','第四人承担后端数据与设备模块'],['SQL 与应用双层约束','异常请求及并发写入都需保护','H2 测试之外仍必须验证 MySQL']]),H('8 扩展顺序'),P('后续优先对接标准化设备接入层、批量查询替代 N+1 查询、数据库迁移和备份恢复演练，再按实测量增加分区或预聚合。Agent 仅在明确需要自然语言查询等功能时评估，需只读权限、可追溯查询与人工确认敏感操作。'),H('9 证据与参考'),P('实现依据为 pom.xml、src/main/java、StatisticsMapper.xml 和 database/01-schema.sql。框架官方参考：docs.spring.io/spring-boot/3.5/；baomidou.com/en/getting-started/install/；vuejs.org/guide/。技术版本核对日期为 2026年9月7日，运行结果以测试日志为准。')]
])

case_data=[
 ('anonymousCannotReadData','FR01','未认证读取','不带会话 GET /elders','HTTP 401'),
 ('realLoginAndLogoutInvalidateSession','FR01','登录注销闭环','获取 CSRF；admin 正确密码登录；GET me；POST logout','登录 200 且无密码字段；me 200；旧 Session 失效'),
 ('invalidPasswordIsRejected','FR01','错误密码','带 CSRF 用 admin 和错误密码登录','HTTP 401'),
 ('csrfAndReadOnlyRoleAreEnforced','FR01','写权限与 CSRF','管理员写入不带 CSRF；ANALYST 带 CSRF 写入；ANALYST 查询','两次写入 403；查询 200'),
 ('roleChangeInvalidatesExistingAuthentication','FR01','会话降权','创建 OPERATOR 保存认证对象；降为 ANALYST；用旧认证访问','HTTP 401'),
 ('lastAdministratorCannotBeDisabled','FR01','最后管理员保护','仅一名有效管理员时尝试停用该账号','409 业务冲突；管理员保留'),
 ('paginationAndMassAssignmentAreRejected','FR03','参数边界','GET elders?size=101；POST elders 提交不允许的 id 字段','均为 400'),
 ('populationIncludesUnknownAndExcludesArchived','FR03','人口未知与归档','创建未知性别出生日期的有效老人；核对总数；归档后重查','初始总数 1 且年龄未知；归档后 0'),
 ('staleElderUpdateIsRejected','FR03','过期版本','对 version=0 档案提交 version=4','409 并提示刷新'),
 ('healthCountsPeopleAndRecordsSeparately','FR04','健康计数','同一老人录入 160 收缩压和 70 心率两条记录','检测 2 次、1 人、异常 1 人'),
 ('healthMissingMetricsAndFutureDataAreRejected','FR04','缺测与未来','所有指标缺省提交；再以明天时间和血氧 99 提交','均拒绝，不入库'),
 ('completionIsAtomicAndCannotBeRepeated','FR05','随访完成','创建一条到期计划；完成；用相同内容再次完成','第二次冲突；只有一条完成记录；完成率 100'),
 ('noPlansHaveNoArtificialCompletionRate','FR05','零计划','空计划集查询随访统计','completionRate 为 null'),
 ('duplicateBindingIsRejected','FR06','重复绑定','设备已绑定该老人，再绑定一次','业务冲突，无重复有效绑定'),
 ('persistentOutsideOnlyCreatesOneAlertAndReturnAllowsAnother','FR07','越界生命周期','500 米圆；连续两个外点；一个内点；再一个外点','事件数 1、1；返回时间写入；最终 2'),
 ('duplicateAndOutOfOrderPointsDoNotCorruptMonitoring','FR07','重传与乱序','外点同内容重传；更早内点；原 eventId 改内容重传','只存两个点；外部状态不回退；改内容 409'),
 ('trajectoryIsSortedAndRestricted','FR08','轨迹顺序与范围','先写较晚点再写较早点；查 1 小时；再查 48 小时','升序返回；48 小时拒绝'),
 ('rebindingPreservesHistoricalOwner','FR06 FR08','换绑历史','解绑老人甲，绑定乙，补传甲绑定期间的定位','历史定位归甲；乙当前定位不被污染'),
 ('heartbeatDoesNotMoveBackwards','FR06','旧心跳','先更新较新心跳，再写更旧心跳','lastSeenAt 不后退'),
 ('productionDoesNotExposeDemoIngestion','FR07','演示边界','非 demo 配置，ADMIN 带 CSRF 调用模拟定位入口','HTTP 404'),
 ('dashboardAndStatisticsContractsWork','FR02','统计响应','认证后访问 dashboard 和三个统计端点','均为 200 且 code=OK'),
 ('distanceHasCorrectScale','FR07','球面距离','计算同点与赤道纬度相差 1 度的距离','同点 0；后一结果 111000 至 111300 米'),
 ('contextLoads','工程','应用上下文','加载默认测试配置与 H2 结构','Spring 上下文启动成功')]
casepages=[[H('1 用例说明'),P('自动化用例以实际 src/test/java 中的方法为准。除登录类特例外，BusinessTests 每项在独立回滚事务中建立一位未知信息老人、一位医生、一只设备和三天前开始的有效绑定。测试管理员密码只用于 H2 测试，不是部署账号密码。'),P('TC01 至 TC23 已于 2026年9月8日执行通过。涉及 HTTP 状态的用例通过 MockMvc 断言；直接业务调用用例断言 Api.Failure 或数据状态，不将其冒充跨网络 HTTP 测试。'),T(['级别','范围','执行方式'],[['自动化功能','认证、统计、历史、围栏边界','mvnw.cmd test'],['MySQL 联调','实际会话、查询、注销','scripts/smoke_http.py'],['浏览器','导航、表单打开、窄屏和只读按钮','scripts/browser-qa.mjs'],['负载','四端点 20 并发 60 秒','scripts/performance_test.py，仅本任务隔离库']])]]
for start in range(0,len(case_data),4):
 page=[H(f'2 自动化用例 {start+1} 至 {min(start+4,len(case_data))}')]
 for idx,(method,fr,title,steps,expected) in enumerate(case_data[start:start+4],start+1):
  page += [H(f'TC{idx:02} {title}'),P(f'需求：{fr}。步骤：{steps}。预期：{expected}。结果：通过。'),P('方法：'+method)]
 casepages.append(page)
casepages.append([H('3 浏览器和人工复核用例'),T(['编号','操作','预期与状态'],[['UI01','逐项进入九个业务导航页','页面加载并可见，无 500 和脚本异常；已通过'],['UI02','人口页打开新增老人对话框','字段及保存取消可见；已通过，不代表所有写表单均端到端提交'],['UI03','390 像素窗口检查人口页','文档无整体横向溢出；已通过'],['UI04','ANALYST 登录人口页','无新增按钮；已通过'],['MAN01','IDEA 选择 JDK 17，右键运行两个测试类','应发现 JUnit 5 并通过；未执行'],['MAN02','轨迹播放 暂停 倍速 拖动进度','时间和点位一致；需团队人工演示复核'],['MAN03','归档 包含设备绑定和待随访的老人','历史保留 绑定关闭 待随访取消；需团队 MySQL 人工复核'],['MAN04','新建专用库进行备份恢复演练','恢复后表数和关键计数一致；未执行']]),P('正式验收时由 A 将人工实际结果、执行人、日期和截图位置写入测试日志。失败项应附复现数据及服务端 requestId，修复后先回归相关用例，再执行核心冒烟。')])
add('05','测试用例','A 组织 B C D 执行','给出真实自动测试映射和人工补充用例',casepages)

add('06','测试计划','A 项目经理兼测试负责人','规定测试范围 资源 顺序 准入准出和缺陷处理',[
 [H('1 目标与范围'),P('本计划用于课程版持续回归与团队验收。验证八项功能、角色隔离、统计口径、设备历史归属和围栏状态，并通过隔离 MySQL 检查实际 SQL。生产安全评估、临床准确性、真实设备联调不属于当前计划。'),T(['阶段','输入','活动和产物'],[['环境确认','代码、pom、专用数据库','记录 Java Maven Node MySQL IDEA 完整版本'],['单体回归','编译成功','JUnit 业务及 MockMvc；Surefire 报告'],['数据库联调','16 表与模拟数据','HTTP 会话、统计、轨迹查询与注销'],['页面验收','可运行 JAR','桌面窄屏、表单、播放、权限与失败状态'],['性能测试','隔离环境可清理专属数据','双十万记录，20 并发 60 秒，端点 P95'],['交付复核','代码文档冻结','部署命令、文件完整性、答辩演示与签收']]),H('2 角色与资源'),P('A 维护计划、用例、缺陷单与验收结论；B 负责认证、健康和随访测试；C 负责 SQL、设备历史、围栏和负载；D 负责页面与浏览器联调。使用虚构数据，不引入个人真实健康或轨迹信息。')],
 [H('3 准入与准出'),T(['环节','准入条件','准出条件'],[['功能测试','依赖下载、测试库可用','全部核心自动测试通过；失败有定位记录'],['接口联调','仅专用库初始化、服务启动成功','16 项冒烟通过；SQL 无方言错误'],['页面验收','前端构建和接口通过','关键页面可使用；阻断与严重缺陷清零'],['性能','记录硬件 数据量 端点 筛选','P95 ≤ 2 秒、错误率 < 1%；恢复原始测试数据'],['交付','文档与代码一致','所有交付文件打开与排版正常；剩余未测项明确列出']]),H('4 缺陷处理'),P('缺陷记录包含编号、版本、前置数据、复现步骤、实际与期望、严重性、负责人及回归结果。阻断为无法登录或无法启动；严重为越权、数据串人、重复业务事实；一般为图表口径或操作错误；轻微为不阻断使用的排版问题。'),P('发现数据库污染、写入到非测试库或账号凭据泄漏时立即停止相关测试并保留最小证据。环境不可用属于未执行，不能计入通过率。并发写入正确性还需专门多连接测试，不以单事务测试替代。'),H('5 相对排期'),P('截止日期未提供，不虚构日历。建议团队验收按 T1 环境与自动回归、T2 MySQL 与页面、T3 性能及缺陷修复、T4 部署与答辩顺序安排；任何失败按影响回退相关阶段。测试日志已记录开发过程实际结果，团队可以在此基础上补充人工验证。')]
])

add('07','项目例会纪要','A 项目经理','保留已确认决策并提供真实例会填写页',[
 [H('1 已确认的沟通决策'),P('本页依据用户在项目沟通中的指令整理，不是虚构的线下例会记录。实际会议时间、参会名单和签字尚未提供，不填入假设的出席事实。'),T(['事项','已确认内容','落实位置'],[['工程方案','使用 Java Spring MVC Spring Boot MyBatis MyBatis-Plus Vue','pom.xml 与前后端代码'],['构建管理','补充 Maven，并增加部署环境说明文档','Maven Wrapper scripts 及第 11 项文档'],['版本调整','Spring Boot 4 降到 3，团队 IDEA 2024 及以下','Boot 3.5.16 与 JUnit 5'],['人员与 Agent','四人：经理测试、后端、后端数据、前端；本版无独立 Agent 岗位','成员分工说明'],['数据与地图','虚构演示数据，圆形围栏和离线坐标画布','DemoSeeder GeoService GeoCanvas'],['重复文档','第 8 项改为系统详细设计说明书','08 系统详细设计说明书'],['继续执行','按现有流程完成开发与交付检查','源代码、测试记录和交付文件']]),H('2 当前需要团队补充'),P('成员真实姓名及学号、完整 IDEA 版本、学校封面模板、实际截止日期、最终验收与答辩时间。A 汇总后更新相关封面与分工；B 补录 JDK 17 和 IDEA 直接测试结果。')],
 [H('3 项目例会填写模板'),T(['字段','待填写内容'],[['会议编号','待填写'],['会议日期 地点或方式','待实际开会后填写'],['主持人 记录人','待填写'],['实际参会及缺席人员','待填写'],['本次目标','待填写']]),H('4 议题与结论'),T(['议题','讨论依据','实际结论'],[['上次行动项检查','待填写','待填写'],['本周功能与缺陷','待填写','待填写'],['风险和范围变更','待填写','待填写']]),H('5 行动项'),T(['行动','负责人','完成日期与验收'],[['补录环境实测','B 姓名待填写','待填写'],['人工演示关键流程','A D 姓名待填写','待填写'],['其他会议新增任务','待填写','待填写']]),P('会议结论确认人及确认时间：待实际确认后填写。保留事实、决策和负责人即可，无需编写与真实项目不符的进度或工时。')]
])

add('08','系统详细设计说明书','B C D 工程师','描述接口规范 输入字段 状态流转 算法与前端实现',[
 [H('1 接口总则'),P('接口前缀 /api/v1，JSON UTF-8。成功返回 {code:"OK",message:"成功",data:...}，列表 data 为 records、total、page、size。Long 值序列化为字符串；时间返回带 +08:00 的 ISO 格式。业务异常含业务码、消息和 requestId；安全过滤器错误仅含业务码和消息。HTTP 状态与错误类型一致。'),T(['方法与路径','说明','权限'],[['GET /auth/csrf','获取 token 和 headerName','匿名可访问'],['POST /auth/login','username password，携带 CSRF','匿名可访问'],['GET /auth/me','当前账号信息','已认证'],['POST /auth/logout','注销 Session，携带 CSRF','已认证'],['GET /dashboard /statistics/*','汇总与图表','三角色'],['POST PUT 业务资源','字段白名单、版本和事务','ADMIN OPERATOR'],['/users 及其子路径','账号列表、维护、密码重置','ADMIN'],['POST /demo/locations /demo/heartbeats','单条模拟事件，仅 demo profile','ADMIN']]),P('首次登录流程为 GET csrf、POST login、再次 GET csrf。前端每次写入使用最新 token；401 通知界面退回登录。不要在浏览器 localStorage 保存密码或 Session 凭据。'),H('2 分页与筛选'),P('分页 page 默认 1，size 默认 20，最大 100。地区与老人、医生 id 为可选查询条件。统计和记录时间默认近 30 天；轨迹必须显式提供 start 与 end。细分异常或状态筛选作用于明细，页面汇总仍按地区与日期显示，使用时注意分母。')],
 [H('3 写入请求字段'),T(['资源','必需字段','可选字段或编辑约束'],[['老人 /elders','code name gender regionId','birthDate phone；编辑必带 version'],['医生 /doctors','code name department enabled','phone；编辑必带 version'],['健康 /health-records','elderId measuredAt；至少一个指标','systolic diastolic heartRate oxygen temperature eventId'],['随访 /followup-plans','elderId doctorId dueAt','完成提交 completedAt content result；取消提交 reason'],['设备 /devices','serialNo model enabled','编辑必带 version；绑定提交 elderId'],['围栏 /geofences','name centerLon centerLat radiusM enabled','编辑必带 version；成员提交 elderIds 数组'],['告警处理 /geofence-alerts/{id}/handle','version note','处理人由服务端当前账号决定'],['账号 /users','username displayName role enabled password','编辑不改 username、不带 password；独立重置接口']]),H('4 字段校验'),P('健康指标输入边界：收缩压 20～300、舒张压 10～200、心率 10～300、血氧 0～100、体温 25～45；同时填写血压时收缩压须大于舒张压。检测时间和随访完成时间不可晚于当前时间。所有这些输入边界均是工程校验，不表示临床正常范围。'),P('教学判定阈值由 health_rule 初始化：收缩压 90～139 mmHg、舒张压 60～89 mmHg、心率 60～100 次/分、血氧 95～100%、体温 36～37.3℃。边界包含，超出即标记；只判断有值指标；保存版本和快照。当前无阈值维护页面，改规则需数据库版本维护与重新验证。')],
 [H('5 定位与围栏算法'),P('单点请求字段 deviceId、eventId、longitude、latitude、recordedAt。经度 -180～180，纬度 -90～90，四舍五入到 7 位小数。查询事件时间对应的 [bound_at,unbound_at) 绑定，写入绑定及老人主键；不存在有效历史绑定则拒收。'),P('距离采用 Haversine：先将纬度和经度差转弧度，a = sin²(Δφ/2) + cos(φ1)cos(φ2)sin²(Δλ/2)，d = 2R atan2(√a,√(1-a))，R 取 6371000 米。d 大于 radiusM 为外部，否则内部。'),T(['原状态与新点','数据库操作','效果'],[['未知或内部收到外点','创建告警，state.active_alert_id 指向新告警','开始一次越界'],['外部收到较新外点','更新最后点，不新增告警','持续越界不重复'],['外部收到较新内点','写 returned_at 与关闭原因，清空活动告警','结束本次越界'],['旧时间定位','入历史表，不推进当前状态','避免迟到点造成假返回'],['围栏修改或停用','关闭原监测并清理状态','下一点依新范围重新判断'],['人工处理','写 handled_by handled_at handling_note','不修改实际 returned_at']]),P('同设备相同 eventId 相同内容返回已存记录；内容冲突返回 409。相同时间使用较大主键决定最新点。状态与事件在同一事务写入，锁顺序保持一致。')],
 [H('6 前端交互设计'),P('App.vue 以 hash 导航管理页面与筛选、分页、弹窗；api.js 统一 fetch、CSRF 与会话过期；Chart.vue 封装 ECharts 并根据 ResizeObserver 调整尺寸；GeoCanvas.vue 将 WGS84 局部投影成坐标画布。地图注明无道路底图，在线状态与定位时刻分别显示。'),T(['页面组','交互与状态'],[['人口 医生 账号','分页列表、维护弹窗、必填校验、版本冲突提示'],['健康 随访','地区日期筛选、汇总图、明细、记录录入或计划完成'],['设备 围栏','坐标选择、详情、绑定及成员配置、告警处理'],['轨迹','老人和日期查询、点数与时间展示、播放暂停倍速和进度'],['通用','加载 失败 空数据；只读用户隐藏写入口；窄屏导航和局部表格滚动']]),H('7 事务与错误恢复'),P('前端遇到 409 提示刷新，不自动覆盖版本；收到 400 提示修正字段；写入发生网络超时时先查询结果再决定是否重试。定位有事件去重，随访有计划唯一结果；普通新增档案依业务编码唯一约束，不能假定所有 POST 都幂等。'),P('源码中的 CareService.enrich 与设备分布存在按条补查，当前数据量已验证，扩大规模时需批量查询并重测。审计只保存动作对象与操作者，不存密码、Session 或整份敏感请求。'),H('8 开发者入口'),P('完整端点目录与请求示例见 docs/API.md。数据字段以 database/01-schema.sql 为准；演示和普通启动参数见 scripts/start.ps1 与项目部署环境说明书。')]
])

add('10','项目组成员分工','A 项目经理','定义四人责任与协作交付边界',[
 [H('1 四人配置'),P('当前八项功能以数据管理、统计和确定性空间规则为主，无须独立 Agent 开发工程师。四人配置采用项目经理兼测试、后端工程师、后端兼数据工程师、前端工程师。姓名和学号待团队补充，下表 A B C D 为责任占位。'),T(['成员','岗位','主要交付'],[['A 姓名待填','项目经理兼测试负责人','需求、计划、例会、用例、日志、进度与答辩汇总'],['B 姓名待填','后端工程师','认证权限、老人医生、健康随访、业务测试、接口说明'],['C 姓名待填','后端兼数据工程师','SQL、模拟数据、设备绑定、围栏轨迹、负载与部署'],['D 姓名待填','前端工程师','Vue 页面、图表、坐标画布、交互联调、演示截图']]),H('2 交叉复核'),T(['交付物','主责','复核'],[['需求与范围变更','A','B C D 评估影响'],['数据库与定位算法','C','B 检查事务和约束'],['认证和业务接口','B','C 检查安全和数据一致性'],['前端交互与图表口径','D','A 验收 B C 核对接口'],['测试与缺陷关闭','A','对应开发负责人修复，其他成员回归'],['部署与答辩','C A','全员各自模块演示']])],
 [H('3 协作流程'),P('接口先明确请求、响应、权限和错误码，再分别实现前后端。每次提交附关联需求和验证结果；SQL 变更新增迁移说明，避免直接修改他人演示库。合并前运行相关测试，阶段完成后运行核心冒烟。'),P('A 维护需求清单与未测项；B C 互审事务和 SQL；D 使用真实接口联调并同步空数据、失败和只读状态。共同修改 App.vue 或数据库结构前先确定责任范围，减少同时覆盖。'),H('4 答辩建议分配'),T(['成员','讲解内容','建议时长'],[['A','背景、目标、范围、测试结论','2 分钟'],['B','Spring Boot 3、权限、统计与随访闭环','3 分钟'],['C','16 表、绑定历史、围栏算法和部署','3 分钟'],['D','八项页面演示、图表、轨迹与交互','4 分钟']]),P('时长为建议安排，并非学校已规定要求。截止日期、正式答辩总时长、成员姓名与贡献工时应由团队按事实补充。'),H('5 Agent 扩展条件'),P('只有新增自然语言统计问答、辅助报告生成等明确需求时，再评估 Agent 工作量与岗位。可先由 C 兼任集成、B 负责权限审计，不应为了岗位名称引入无法验收的功能。Agent 输出须展示数据来源，且不替代临床判断。')]
])

add('11','项目部署环境说明书','C 部署负责人 B 协助','说明环境 凭据 初始化 构建 启动 验证及故障处理',[
 [H('1 部署结论与环境矩阵'),P('交付形态为 Vue 静态资源内置的 Spring Boot 可执行 JAR，加一套 MySQL 数据库。后端 Maven 构建，前端 npm 构建。首次部署需建专用空库与表结构，再设置数据库凭据和初始管理员密码。'),T(['组件','团队基线','本机验证'],[['操作系统','Windows 10 或 11 开发；其他系统另测','Windows 11'],['Java','团队建议 JDK 17，pom release 17','Zulu 25.0.1 编译 测试 运行通过'],['IDEA','2024 系列及以下按完整版本测试','版本未提供，内置 JUnit 未实测'],['Maven','Wrapper 3.9.16','编译 测试 打包通过'],['Node 与前端','Node 22.12 以上的 22 系列或满足 Vite 7 要求','Node 22.22.0；Vite 7.3.6 构建通过'],['MySQL','8.0.16 以上 8.0 系列；utf8mb4；上海时区','8.0.46 独立进程，16 表通过'],['浏览器','现代 Chrome 或 Edge','Chrome 152.0.7977.83'],['端口','应用 8080 数据库 3306 可配置','应用 18080 隔离 MySQL 13316']]),P('资源建议为 4 核 CPU、8 GB 内存和 SSD；这只是小型教学部署建议，不是经测量的最低配置。本次性能 CPU 为 i7-13650HX，结果详见测试日志。')],
 [H('2 首次初始化'),P('在 MySQL 管理客户端执行 database/00-create-database.sql，再选择 smart_care 执行 database/01-schema.sql。脚本不含 DROP，非空库不要重复运行。为应用创建专用账号并仅授权 smart_care 库；管理密码不要写进代码或命令行参数。'),P('示例 SQL：CREATE USER \'smart_care\'@\'127.0.0.1\' IDENTIFIED BY \'替换为自己的随机密码\'; GRANT SELECT, INSERT, UPDATE, DELETE ON smart_care.* TO \'smart_care\'@\'127.0.0.1\'; 建表使用管理账号，应用运行不需要 DDL 权限。'),P('使用 mysql -h 127.0.0.1 -P 3306 -u 管理账号 -p，交互输入密码；再用 SOURCE D:/项目路径/database/00-create-database.sql; 与 SOURCE D:/项目路径/database/01-schema.sql;。路径避免空格或通过客户端选择脚本，字符集设 utf8mb4。'),H('3 必需环境变量'),T(['变量','用途'],[['DB_URL','jdbc:mysql://127.0.0.1:3306/smart_care?useUnicode=true&characterEncoding=UTF-8&connectionTimeZone=Asia/Shanghai'],['DB_USERNAME DB_PASSWORD','专用应用数据库账号和密码'],['APP_ADMIN_PASSWORD','首次空账号库创建 admin，至少 12 字符、不超过 72 UTF-8 字节'],['SERVER_PORT SERVER_ADDRESS','默认 8080 与 127.0.0.1'],['SPRING_PROFILES_ACTIVE','demo 为模拟入口；prod 为 HTTPS Cookie 配置'],['APP_SEED_DEMO','仅 demo 配合 true，空老人库填充虚构数据']]),P('.env.example 只用于说明，Spring Boot 不自动读取 .env。可在当前 PowerShell 设置 $env:变量，或把私有 JSON 路径交给 scripts/start.ps1 -EnvironmentFile；私有文件不得提交 Git。')],
 [H('4 构建与运行'),P('在项目根目录 PowerShell 执行 .\\scripts\\build.ps1。脚本依次 npm ci、npm run build、复制静态资源、mvnw.cmd package；默认执行测试。首次构建需要访问 Maven 和 npm 仓库，后续受本机缓存影响。生成 target/dev-0.0.1-SNAPSHOT.jar。'),P('构建前先停止正在运行的同一 JAR。Windows 会锁定文件，未停止可能出现 repackage 无法重命名。前台启动可按 Ctrl+C；不要按进程名批量结束其他 Java 服务。'),P('演示启动：.\\scripts\\start.ps1 -Profile demo -Port 18080 -EnvironmentFile .\\.local\\demo-env.json。该命令在前台运行，浏览器打开 http://127.0.0.1:18080。账号 admin，密码来自自己设置的 APP_ADMIN_PASSWORD。演示空库还会创建 operator 与 analyst，初始演示密码相同，正式使用须分别更改。'),P('普通启动：.\\scripts\\start.ps1 -Profile local -Port 8080 -EnvironmentFile 私有JSON路径。已有账号库不会因修改 APP_ADMIN_PASSWORD 而重置密码；请通过管理员界面重置。'),H('5 演示数据选择'),P('推荐 Java DemoSeeder：空库首次 demo 启动生成 120 老人、6 医生、1080 健康记录、120 随访计划、18 设备、648 定位、1 围栏。时间相对启动时生成，便于答辩。database/02-demo-data.sql 为独立固定日期 SQL 种子，二者只能选一种，不能向非空业务库重复导入。'),P('当前任务已经初始化的隔离库位于临时目录，端口 13316；本机访问参数见 .local/demo-access.txt。该临时库用于本次验证，重启机器或清理临时目录后需按上述步骤重新建立持久专用库。')],
 [H('6 开发与测试'),P('前后端分开开发：后端 .\\mvnw.cmd spring-boot:run，前端进入 frontend 执行 npm ci、npm run dev。Vite 默认将 /api 代理到 8080；如果后端改端口，同步修改 vite.config.js 的代理。交付 JAR 采用同源，无需 Vite 服务。'),P('独立测试：.\\mvnw.cmd test。测试使用 H2 与测试配置，不需要 MySQL 密码。IDEA 选择 Project SDK 17 和 Maven Runner JRE 17，刷新 Maven 后对 BusinessTests 或 DevApplicationTests 右键运行。使用 Boot BOM 管理的 JUnit 5 API、Engine 与 Launcher，避免手工加入 JUnit 6。'),P('实际 HTTP：用 Python 执行 scripts/smoke_http.py --base-url http://127.0.0.1:18080 --env-file 私有JSON路径。浏览器验证脚本需要 Playwright 与 Chrome；当前工具运行时由 RUNTIME_NODE_MODULES 指定包目录，此环境为本次开发自动化依赖，不是应用运行依赖。'),H('7 常见故障'),T(['现象','处理'],[['启动提示缺少 APP_ADMIN_PASSWORD','仅首次空账号库必须设置符合长度的密码'],['Access denied 或连接失败','核对 host 端口 库名 账号授权与服务状态'],['8080 已占用','更换 SERVER_PORT，不停止无关服务'],['JAR 无法重命名','先关闭本项目正在运行的 JAR，再 build'],['prod 登录后会话不保持','Secure Cookie 要求 HTTPS；本地 HTTP 用 local 或 demo'],['图表无数据或设备全离线','检查时间范围、模拟数据日期；5 分钟无新心跳转离线'],['Maven 依赖解析失败','检查仓库网络与 settings.xml，不混改 JUnit 单包版本']])],
 [H('8 运维与回滚'),P('数据库业务日期依赖 MySQL 会话的 CURRENT_DATE，数据库系统时区应设为 +08:00，并与应用 Asia/Shanghai 保持一致。正式入口通过 HTTPS 反向代理访问 prod 配置；证书、域名、防火墙与公共发布尚未执行，必须按最终目标环境确认。'),P('备份示例：mysqldump --single-transaction --routines --triggers -h 127.0.0.1 -P 3306 -u 备份账号 -p --result-file=backup.sql smart_care。密码交互输入。先在独立恢复库验证表结构、行数和关键业务查询，再考虑业务切换；不得直接覆盖当前数据库。'),P('发布前保留旧 JAR 与配置，记录 SQL 版本和备份文件。仅代码回滚且结构兼容时可停止服务换回旧 JAR；有结构变化时需单独恢复方案。当前没有自动迁移、自动备份或集群会话共享。'),H('9 部署验收记录'),T(['检查项','本次状态'],[['16 表结构执行','隔离 MySQL 通过'],['静态资源打包与 JAR 启动','本机通过'],['Maven JUnit 与 HTTP 冒烟','23 项与 16 项通过'],['前端页面和只读权限','12 项浏览器检查通过'],['压力测试','2664 次请求，0 错误，P95 达标'],['JDK 17 IDEA 直接运行','待团队实测'],['持久环境 HTTPS 备份恢复','待目标部署确认后执行']]),P('环境复核人、正式部署路径、目标服务器配置和签收日期：待团队实际部署后填写。')]
])

def setfont(run,size=None,bold=None,color=None):
 run.font.name='Georgia';run._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'Microsoft YaHei')
 if size:run.font.size=Pt(size)
 if bold is not None:run.bold=bold
 if color:run.font.color.rgb=RGBColor.from_string(color)
def paragraph(doc,text,style='normal',size=11):
 p=doc.add_paragraph(style=style);p.paragraph_format.space_after=Pt(7);p.paragraph_format.line_spacing=1.18
 grid=OxmlElement('w:snapToGrid');grid.set(qn('w:val'),'0');p._p.get_or_add_pPr().append(grid)
 for lineidx,line in enumerate(text.split('\n')):
  if lineidx:p.add_run().add_break()
  setfont(p.add_run(line),size)
 return p
def table(doc,heads,rows):
 elem=deepcopy(table_patterns[len(heads)]);doc._body._body.insert(-1,elem);tbl=Table(elem,doc._body)
 # Keep template table geometry and cell styling, adapt row count to actual data.
 while len(tbl.rows)>2:tbl._tbl.remove(tbl.rows[-1]._tr)
 body=deepcopy(tbl.rows[1]._tr)
 while len(tbl.rows)<len(rows)+1:tbl._tbl.append(deepcopy(body))
 for ri,row in enumerate([heads]+rows):
  tr=tbl.rows[ri]._tr;props=tr.get_or_add_trPr()
  for old in list(props):
   if old.tag==qn('w:trHeight'):props.remove(old)
  props.append(OxmlElement('w:cantSplit'))
  if ri==0:props.append(OxmlElement('w:tblHeader'))
  for ci,value in enumerate(row):
   cell=tbl.cell(ri,ci);cell.text=str(value)
   margins=cell._tc.get_or_add_tcPr().find(qn('w:tcMar'))
   if margins is None:margins=OxmlElement('w:tcMar');cell._tc.get_or_add_tcPr().append(margins)
   for edge in ('top','bottom','left','right'):
    node=OxmlElement('w:'+edge);node.set(qn('w:w'),'45' if edge in ('top','bottom') else '85');node.set(qn('w:type'),'dxa');margins.append(node)
   for p in cell.paragraphs:
    p.paragraph_format.space_before=Pt(1);p.paragraph_format.space_after=Pt(1);p.paragraph_format.line_spacing=1.12
    grid=OxmlElement('w:snapToGrid');grid.set(qn('w:val'),'0');p._p.get_or_add_pPr().append(grid)
    for r in p.runs:setfont(r,9.5,ri==0,'FFFFFF' if ri==0 else '263C50')
  if ri==0:
   for cell in tbl.rows[ri].cells:
    shd=cell._tc.get_or_add_tcPr().find(qn('w:shd'))
    if shd is None:shd=OxmlElement('w:shd');cell._tc.get_or_add_tcPr().append(shd)
    shd.set(qn('w:fill'),'0B2C48')
 paragraph(doc,'',size=3).paragraph_format.space_after=Pt(0)
 return tbl

def build(spec):
 doc=Document(REF)
 for el in list(doc._body._body):
  if el.tag!=qn('w:sectPr'):doc._body._body.remove(el)
 for sec in doc.sections:
  for p in sec.footer.paragraphs:p.text=''
  p=sec.footer.paragraphs[0];p.alignment=1;r=p.add_run(PROJECT+'  /  V1.0  /  ');setfont(r,8,color='5E758A')
  fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');p._p.append(fld)
 # Preserve the reference's generous cover and two-part title hierarchy.
 p=paragraph(doc,'',size=10);p.paragraph_format.space_after=Pt(125)
 p=paragraph(doc,'智慧医养大数据\n公共服务平台','Title',25);p.paragraph_format.space_after=Pt(17)
 for r in p.runs:setfont(r,25,False,'5E758A')
 p=paragraph(doc,spec['title'],'Title',29);p.paragraph_format.space_after=Pt(90)
 for r in p.runs:setfont(r,29,True,'000000')
 table(doc,['版本与状态','责任人','更新日期'],[['V1.0 课程版交付',spec['owner'],DATE]])
 table(doc,['文档信息','内容'],[['系统',PROJECT],['范围',spec['scope']],['读者','项目组成员 指导教师 验收人员'],['说明','姓名与实际签收信息由团队补充；测试结论按事实记录']])
 for page in spec['pages']:
  doc.add_page_break()
  for block in page:
   if block[0]=='h':
    p=paragraph(doc,block[1],'Heading 1',15);p.paragraph_format.space_before=Pt(10);p.paragraph_format.space_after=Pt(8);p.paragraph_format.keep_with_next=True
    for r in p.runs:setfont(r,15,True,'0B2C48')
   elif block[0]=='p':paragraph(doc,block[1])
   elif block[0]=='t':table(doc,block[1],block[2])
   elif block[0]=='img':doc.add_picture(str(ROOT/block[1]),width=Inches(6.9));paragraph(doc,block[2],size=9)
 doc.core_properties.title=PROJECT+' '+spec['title'];doc.core_properties.author='智慧医养项目组';doc.core_properties.subject=spec['scope'];doc.core_properties.comments='Generated from verified repository sources. Team review pending.'
 out=OUT/(spec['number']+' '+spec['title']+'.docx');doc.save(out)
 md=['# '+spec['title'],'版本 V1.0　'+DATE]
 for page in spec['pages']:
  for b in page:
   if b[0]=='h':md+=['','## '+b[1]]
   elif b[0]=='p':md+=['',b[1]]
   elif b[0]=='t':md+=['','| '+' | '.join(b[1])+' |','| '+' | '.join(['---']*len(b[1]))+' |']+['| '+' | '.join(map(str,r))+' |' for r in b[2]]
 (SOURCE/(spec['number']+'-'+spec['title']+'.md')).write_text('\n'.join(md)+'\n',encoding='utf-8')
 print(out.name)
import sys
for spec in docs:
 if len(sys.argv)==1 or spec['number'] in sys.argv[1:]:build(spec)
assert hashlib.sha256(REF.read_bytes()).hexdigest()=='13504f6c221a42c1726460a9e865e563355539ff97d702d6c9b2267b4b261d76'
(ROOT/'docs/.qa/document-manifest.json').write_text(json.dumps([{'file':s['number']+' '+s['title']+'.docx','plannedPages':1+len(s['pages'])} for s in docs],ensure_ascii=False,indent=2),encoding='utf-8')
