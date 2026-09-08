from pathlib import Path
from zipfile import ZipFile
from xml.etree import ElementTree as ET
from docx import Document
import json,hashlib
root=Path(__file__).resolve().parents[1]
templates=Path('C:/Users/29569/.codex/plugins/cache/openai-curated-remote/openai-templates/0.1.1/skills')
docpath=templates/'artifact-template-system-design/assets/reference.docx'
doc=Document(docpath)
output={'reference':str(docpath),'sha256':hashlib.sha256(docpath.read_bytes()).hexdigest(),'sections':[], 'paragraphs':[], 'styles':[], 'tables':[]}
for s in doc.sections:output['sections'].append({k:str(getattr(s,k)) for k in ['page_width','page_height','top_margin','bottom_margin','left_margin','right_margin','header_distance','footer_distance']})
for i,p in enumerate(doc.paragraphs):
    if p.text:output['paragraphs'].append({'index':i,'style':p.style.name,'text':p.text})
for s in doc.styles:
    if s.name in ['Title','Subtitle','Heading 1','Heading 2','Normal','Caption']:
        output['styles'].append({'name':s.name,'font':s.font.name,'size':str(s.font.size),'xml':s.element.xml})
for table in doc.tables:output['tables'].append([[c.text for c in r.cells] for r in table.rows])
qa=root/'docs/.qa';qa.mkdir(parents=True,exist_ok=True)
(qa/'doc-template.json').write_text(json.dumps(output,ensure_ascii=False,indent=2),encoding='utf-8')
print('DOCX',json.dumps({k:output[k] for k in ['sections','paragraphs','tables']},ensure_ascii=False))
ns={'a':'http://schemas.openxmlformats.org/drawingml/2006/main','p':'http://schemas.openxmlformats.org/presentationml/2006/main'}
pptpath=templates/'artifact-template-team-alignment/assets/reference.pptx'
with ZipFile(pptpath) as z:
    print('PPTX size',ET.fromstring(z.read('ppt/presentation.xml')).find('p:sldSz',ns).attrib)
    slides=sorted([n for n in z.namelist() if n.startswith('ppt/slides/slide') and n.endswith('.xml')],key=lambda x:int(x.split('slide')[-1].split('.')[0]))
    result=[]
    for name in slides:
        slide=ET.fromstring(z.read(name));texts=[''.join(s.itertext()) for s in slide.findall('.//a:t',ns)]
        result.append({'part':name,'texts':texts});print(name, json.dumps(texts,ensure_ascii=False))
    (qa/'ppt-template.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8')
