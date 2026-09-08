"""Structural QA for generated Word deliverables when no Office renderer exists."""
import json
import zipfile
from datetime import datetime, timezone
from pathlib import Path
from docx import Document

root = Path(__file__).resolve().parents[1]
files = sorted((root / "docs/deliverables").glob("*.docx"))
results = []
for path in files:
    with zipfile.ZipFile(path) as package:
        broken_member = package.testzip()
        names = set(package.namelist())
        package_text = b"\n".join(package.read(name) for name in names if name.startswith("word/") and name.endswith(".xml"))
    document = Document(path)
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    result = {
        "file": path.name,
        "zipValid": broken_member is None,
        "hasMainDocument": "word/document.xml" in names,
        "versionV11Present": b"V1.1" in package_text,
        "paragraphs": len(document.paragraphs),
        "tables": len(document.tables),
        "nonEmptyParagraphs": sum(bool(paragraph.text.strip()) for paragraph in document.paragraphs),
    }
    result["passed"] = all((result["zipValid"], result["hasMainDocument"], result["versionV11Present"], result["nonEmptyParagraphs"] >= 10))
    results.append(result)

report = {
    "runAt": datetime.now(timezone.utc).isoformat(),
    "expectedDocuments": 10,
    "actualDocuments": len(files),
    "allPassed": len(files) == 10 and all(item["passed"] for item in results),
    "visualRender": {
        "passed": False,
        "reason": "LibreOffice soffice.exe and Microsoft Word are unavailable in this execution environment",
    },
    "documents": results,
}
output = root / "docs/evidence/document-v11-structural-qa.json"
output.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
print(json.dumps({"documents": len(files), "allPassed": report["allPassed"], "report": str(output)}, ensure_ascii=False))
raise SystemExit(0 if report["allPassed"] else 1)
